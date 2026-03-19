"""
RAG Pipeline: Retrieval-Augmented Generation.

Based on Ch 8: Understanding agent memory and knowledge.

RAG augments an agent's prompts with relevant context retrieved from
a knowledge base using vector embeddings and similarity search. This
is how agents access information beyond what is in their training data.

Pipeline:
  1. INGEST   - Load documents, chunk them, compute embeddings
  2. INDEX    - Store chunks + embeddings in vector store
  3. RETRIEVE - Given a query, find the most relevant chunks
  4. AUGMENT  - Inject retrieved context into the agent's prompt
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import structlog
from pydantic import BaseModel

logger = structlog.get_logger()


class Document(BaseModel):
    """A document chunk with metadata."""

    content: str
    metadata: dict[str, Any] = {}
    doc_id: str = ""
    chunk_index: int = 0


class RetrievalResult(BaseModel):
    """A single retrieval result with relevance score."""

    document: Document
    score: float
    rank: int


class RAGPipeline:
    """
    Full RAG pipeline: ingest, index, retrieve, augment.

    Usage:
        rag = RAGPipeline(persist_dir="./data/embeddings")
        rag.ingest_directory("./data/knowledge_base")
        results = rag.retrieve("What is federated learning?", top_k=5)
        context = rag.format_context(results)
    """

    def __init__(
        self,
        persist_dir: str = "./data/embeddings",
        collection_name: str = "knowledge_base",
        embedding_model: str = "all-MiniLM-L6-v2",
        chunk_size: int = 500,
        chunk_overlap: int = 50,
    ):
        self.persist_dir = persist_dir
        self.collection_name = collection_name
        self.embedding_model = embedding_model
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self._collection = None

    def _get_collection(self):
        """Lazy-load ChromaDB collection."""
        if self._collection is None:
            import chromadb

            client = chromadb.PersistentClient(path=self.persist_dir)
            self._collection = client.get_or_create_collection(
                name=self.collection_name,
                metadata={"hnsw:space": "cosine"},
            )
        return self._collection

    def chunk_text(self, text: str, doc_id: str = "") -> list[Document]:
        """
        Split text into overlapping chunks for embedding.

        Uses a simple character-based chunking strategy with overlap
        to preserve context across chunk boundaries.
        """
        chunks = []
        start = 0
        chunk_index = 0

        while start < len(text):
            end = start + self.chunk_size

            # Try to break at sentence boundary
            if end < len(text):
                last_period = text.rfind(".", start, end)
                last_newline = text.rfind("\n", start, end)
                break_point = max(last_period, last_newline)
                if break_point > start:
                    end = break_point + 1

            chunk_text = text[start:end].strip()
            if chunk_text:
                chunks.append(
                    Document(
                        content=chunk_text,
                        doc_id=doc_id,
                        chunk_index=chunk_index,
                        metadata={"source": doc_id, "chunk": chunk_index},
                    )
                )
                chunk_index += 1

            start = end - self.chunk_overlap

        return chunks

    def ingest_file(self, file_path: str | Path) -> int:
        """Ingest a single file into the knowledge base. Returns chunk count."""
        path = Path(file_path)
        text = ""

        if path.suffix == ".txt" or path.suffix == ".md":
            text = path.read_text(encoding="utf-8")
        elif path.suffix == ".pdf":
            try:
                from pypdf import PdfReader

                reader = PdfReader(str(path))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
            except ImportError:
                logger.warning("pypdf not installed, skipping PDF", file=str(path))
                return 0
        elif path.suffix == ".docx":
            try:
                import docx

                doc = docx.Document(str(path))
                text = "\n".join(p.text for p in doc.paragraphs)
            except ImportError:
                logger.warning("python-docx not installed, skipping DOCX", file=str(path))
                return 0
        else:
            # Try reading as plain text
            try:
                text = path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                logger.warning("cannot_read_file", file=str(path))
                return 0

        if not text.strip():
            return 0

        chunks = self.chunk_text(text, doc_id=str(path))
        collection = self._get_collection()

        collection.add(
            ids=[f"{path.stem}_{c.chunk_index}" for c in chunks],
            documents=[c.content for c in chunks],
            metadatas=[c.metadata for c in chunks],
        )

        logger.info("file_ingested", file=str(path), chunks=len(chunks))
        return len(chunks)

    def ingest_directory(self, directory: str | Path) -> int:
        """Ingest all supported files in a directory. Returns total chunk count."""
        path = Path(directory)
        total = 0
        supported = {".txt", ".md", ".pdf", ".docx", ".py", ".yaml", ".json"}

        for file_path in sorted(path.rglob("*")):
            if file_path.is_file() and file_path.suffix in supported:
                total += self.ingest_file(file_path)

        logger.info("directory_ingested", directory=str(path), total_chunks=total)
        return total

    def retrieve(
        self,
        query: str,
        top_k: int = 5,
        where: dict[str, Any] | None = None,
    ) -> list[RetrievalResult]:
        """
        Retrieve the most relevant document chunks for a query.

        Args:
            query: The search query.
            top_k: Number of results to return.
            where: Optional metadata filter.

        Returns:
            List of RetrievalResult sorted by relevance.
        """
        collection = self._get_collection()
        kwargs: dict[str, Any] = {
            "query_texts": [query],
            "n_results": top_k,
        }
        if where:
            kwargs["where"] = where

        results = collection.query(**kwargs)

        retrieval_results = []
        if results and results["documents"]:
            for i, (doc_text, metadata, distance) in enumerate(
                zip(
                    results["documents"][0],
                    results["metadatas"][0],
                    results["distances"][0],
                )
            ):
                retrieval_results.append(
                    RetrievalResult(
                        document=Document(content=doc_text, metadata=metadata),
                        score=1 - distance,  # Convert distance to similarity
                        rank=i + 1,
                    )
                )

        return retrieval_results

    def format_context(
        self,
        results: list[RetrievalResult],
        max_tokens: int = 2000,
    ) -> str:
        """
        Format retrieval results into context text for prompt injection.

        This is the AUGMENT step: taking retrieved documents and
        formatting them so they can be prepended to an agent's prompt.
        """
        context_parts = []
        estimated_tokens = 0

        for result in results:
            chunk_text = (
                f"[Source: {result.document.metadata.get('source', 'unknown')} | "
                f"Relevance: {result.score:.2f}]\n"
                f"{result.document.content}\n"
            )
            # Rough token estimate: ~4 chars per token
            chunk_tokens = len(chunk_text) // 4
            if estimated_tokens + chunk_tokens > max_tokens:
                break

            context_parts.append(chunk_text)
            estimated_tokens += chunk_tokens

        if not context_parts:
            return ""

        header = "--- Retrieved Context ---\n"
        footer = "\n--- End Context ---"
        return header + "\n".join(context_parts) + footer
